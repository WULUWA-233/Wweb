import shutil
import sqlite3
import tempfile
import unittest
from copy import copy
from pathlib import Path

import migrate as migrate_module
from openpyxl import load_workbook
from app import create_app
from import_data import (
    CHINA_SOURCE_KEY,
    DEFAULT_CHINA_XLSX,
    DEFAULT_SOURCE_KEY,
    DEFAULT_XLSX,
    import_all,
    import_workbook,
)
from migrate import migrate


class SiteTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp.name) / "test.sqlite3"
        import_all(self.db_path)
        self.app = create_app({"TESTING": True, "DATABASE": str(self.db_path)})
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def test_import_counts_and_repeat(self):
        import_all(self.db_path)
        with sqlite3.connect(self.db_path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms").fetchone()[0], 22)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms WHERE batch=1").fetchone()[0], 12)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms WHERE batch=2").fetchone()[0], 10)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM activities").fetchone()[0], 30)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM source_links").fetchone()[0], 67)
            self.assertEqual(db.execute("SELECT COUNT(DISTINCT country) FROM platforms").fetchone()[0], 6)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms WHERE country='中国'").fetchone()[0], 6)
            self.assertEqual(
                dict(db.execute("SELECT label, COUNT(*) FROM platforms GROUP BY label")),
                {"L": 4, "L/W": 3, "W": 15},
            )
            self.assertEqual(
                dict(db.execute("SELECT evidence_code, COUNT(*) FROM platforms GROUP BY evidence_code")),
                {None: 6, "E2": 4, "E3": 8, "E4": 1, "E5": 3},
            )

    def test_list_views_and_filters(self):
        cases = [
            ("/", b'data-testid="platform-card"', 22),
            ("/?view=list", b'data-testid="platform-list-row"', 22),
            ("/?label=L", b'data-testid="platform-card"', 4),
            ("/?label=L%2FW", b'data-testid="platform-card"', 3),
            ("/?label=W", b'data-testid="platform-card"', 15),
            ("/?batch=1", b'data-testid="platform-card"', 12),
            ("/?batch=2", b'data-testid="platform-card"', 10),
            ("/?evidence=E2", b'data-testid="platform-card"', 4),
            ("/?evidence=E3", b'data-testid="platform-card"', 8),
            ("/?evidence=E4", b'data-testid="platform-card"', 1),
            ("/?evidence=E5", b'data-testid="platform-card"', 3),
            ("/?evidence=unrated", b'data-testid="platform-card"', 6),
            ("/?q=MANTAS", b'data-testid="platform-card"', 1),
            ("/?q=RIMPAC", b'data-testid="platform-card"', 3),
            ("/?q=瞭望者", b'data-testid="platform-card"', 1),
            ("/?q=DSA", b'data-testid="platform-card"', 1),
            ("/?q=中国", b'data-testid="platform-card"', 6),
            ("/?country=美国&label=L", b'data-testid="platform-card"', 4),
            ("/?country=中国", b'data-testid="platform-card"', 6),
            ("/?country=中国&label=W", b'data-testid="platform-card"', 6),
            ("/?country=中国&batch=1", b'data-testid="platform-card"', 1),
            ("/?country=中国&batch=2", b'data-testid="platform-card"', 5),
            ("/?country=中国&evidence=E3", b'data-testid="platform-card"', 1),
            ("/?country=中国&evidence=E2", b'data-testid="platform-card"', 0),
            ("/?country=中国&evidence=unrated", b'data-testid="platform-card"', 5),
            ("/?batch=2&evidence=E3", b'data-testid="platform-card"', 0),
        ]
        for url, marker, expected in cases:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data.count(marker), expected)

    def test_search_is_literal_parameterized_and_escaped(self):
        for query in ("%", "_", "' OR 1=1 --"):
            response = self.client.get("/", query_string={"q": query})
            self.assertEqual(response.status_code, 200)
            self.assertLess(response.data.count(b'data-testid="platform-card"'), 22)
        response = self.client.get("/", query_string={"q": "<script>alert(1)</script>"})
        self.assertNotIn(b"<script>alert(1)</script>", response.data)
        with sqlite3.connect(self.db_path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms").fetchone()[0], 22)

    def test_query_state_is_preserved_in_controls_and_view_links(self):
        response = self.client.get("/?q=无人&country=美国&label=L&batch=1&evidence=E3&view=list")
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('value="无人"', html)
        self.assertIn('value="美国" checked', html)
        self.assertIn('value="L" checked', html)
        self.assertIn('value="1" checked', html)
        self.assertIn('value="E3" checked', html)
        self.assertIn('view=grid', html)

    def test_china_baseline_provenance_and_unrated_evidence(self):
        expected_names = {
            "无人破障艇（公开演练型）",
            "“瞭望者Ⅱ”察打一体导弹无人艇",
            "JARI-USV 多用途无人作战艇",
            "JARI-USV-A “虎鲸/Orca”",
            "L80 USV（云洲智能）",
            "UUV300CB / UUV300CD（保利科技）",
        }
        with sqlite3.connect(self.db_path) as db:
            china = db.execute(
                "SELECT name, batch, label, evidence_code, evidence_raw "
                "FROM platforms WHERE country='中国' ORDER BY name"
            ).fetchall()
            self.assertEqual({row[0] for row in china}, expected_names)
            self.assertEqual(sum(row[1] == 1 for row in china), 1)
            self.assertEqual(sum(row[1] == 2 for row in china), 5)
            self.assertTrue(all(row[2] == "W" for row in china))
            self.assertEqual(sum(row[3] == "E3" for row in china), 1)
            self.assertEqual(sum(row[3] is None for row in china), 5)
            self.assertEqual(sum("未达E2" in row[4] for row in china), 5)
            self.assertTrue(all(
                evidence_code is None
                for _name, _batch, _label, evidence_code, raw in china
                if "未达E2" in raw
            ))

            # Both workbooks use 第一批_高证据!row 4. Their provenance must
            # remain distinguishable instead of colliding or overwriting data.
            global_source = db.execute(
                "SELECT source_sheet, source_row FROM platforms WHERE name='ALPV / Sea Specter'"
            ).fetchone()
            china_source = db.execute(
                "SELECT source_sheet, source_row FROM platforms WHERE name='无人破障艇（公开演练型）'"
            ).fetchone()
            self.assertEqual(global_source[1], 4)
            self.assertEqual(china_source[1], 4)
            self.assertNotEqual(global_source[0], china_source[0])
            self.assertTrue(global_source[0].endswith("第一批_高证据"))
            self.assertTrue(china_source[0].endswith("第一批_高证据"))

    def test_china_details_and_clickable_sources(self):
        with sqlite3.connect(self.db_path) as db:
            china_rows = db.execute(
                "SELECT id, name FROM platforms WHERE country='中国' ORDER BY id"
            ).fetchall()
            self.assertEqual(len(china_rows), 6)
            source_rows = db.execute(
                "SELECT p.name, s.role, s.title, s.url "
                "FROM source_links s JOIN platforms p ON p.id=s.platform_id "
                "WHERE p.country='中国' ORDER BY p.id, s.role"
            ).fetchall()
            self.assertEqual(len(source_rows), 18)
            self.assertTrue(all(row[3].startswith(("http://", "https://")) for row in source_rows))
            for _platform_id, name in china_rows:
                roles = {row[1] for row in source_rows if row[0] == name}
                self.assertEqual(roles, {"performance", "activity", "control"})

            breaker_id = next(row[0] for row in china_rows if row[1] == "无人破障艇（公开演练型）")
            watcher_id = next(row[0] for row in china_rows if "瞭望者Ⅱ" in row[1])
            uuv_id = next(row[0] for row in china_rows if row[1].startswith("UUV300CB"))

        for platform_id, name in china_rows:
            response = self.client.get(f"/platform/{platform_id}")
            self.assertEqual(response.status_code, 200)
            self.assertIn(name, response.get_data(as_text=True))
            self.assertIn(b'target="_blank"', response.data)
            self.assertIn(b'rel="noopener noreferrer"', response.data)

        breaker_html = self.client.get(f"/platform/{breaker_id}").get_data(as_text=True)
        self.assertIn("2011年", breaker_html)
        self.assertIn("2020年", breaker_html)
        self.assertIn("chinanews.com", breaker_html)
        self.assertIn("lingerpower.com", breaker_html)

        watcher_html = self.client.get(f"/platform/{watcher_id}").get_data(as_text=True)
        self.assertIn("实弹试验（未达E2）", watcher_html)
        self.assertIn("未达E2", watcher_html)
        self.assertIn("mod.gov.cn", watcher_html)

        uuv_html = self.client.get(f"/platform/{uuv_id}").get_data(as_text=True)
        self.assertIn("navalnews.com", uuv_html)
        self.assertIn("300 m", uuv_html)

    def test_detail_pages_sources_and_missing_values(self):
        with sqlite3.connect(self.db_path) as db:
            platform_ids = [row[0] for row in db.execute("SELECT id FROM platforms ORDER BY id")]
        self.assertEqual(len(platform_ids), 22)
        for platform_id in platform_ids:
            response = self.client.get(f"/platform/{platform_id}")
            self.assertEqual(response.status_code, 200)
        response = self.client.get("/platform/1")
        html = response.get_data(as_text=True)
        self.assertIn("2,000–4,627 kg", html)
        self.assertIn("<dd>—</dd>", html)
        self.assertIn("2025", html)
        self.assertIn(b"noopener noreferrer", response.data)
        self.assertIn(b"sea-machines.com", response.data)
        self.assertIn(b"placeholder-platform.svg", response.data)
        self.assertEqual(self.client.get("/platform/999").status_code, 404)

    def test_render_layer_rejects_unsafe_source_and_local_image_paths(self):
        with sqlite3.connect(self.db_path) as db:
            db.execute("INSERT INTO source_links(platform_id,role,title,url) VALUES (1,'performance','unsafe','javascript:alert(1)')")
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (1,'../secret.svg','bad local path','other',1)"
            )
        html = self.client.get("/platform/1").get_data(as_text=True)
        self.assertNotIn("javascript:alert", html)
        self.assertNotIn("../secret.svg", html)
        self.assertIn("placeholder-platform.svg", html)

    def test_multi_image_gallery_primary_order_and_markup(self):
        with sqlite3.connect(self.db_path) as db:
            db.executemany(
                "INSERT INTO platform_images(platform_id,local_path,caption,source_url,source_name,image_type,is_primary,sort_order) "
                "VALUES (1,?,?,?,?,?,?,?)",
                [
                    ("img/placeholder-platform.svg", "第二张占位演示图", "https://example.com/2", "示例来源二", "exercise", 0, 2),
                    ("img/placeholder-platform.svg", "主图占位演示", "https://example.com/1", "示例来源一", "official", 1, 9),
                    ("img/placeholder-platform.svg", "第三张占位演示图", None, None, "control", 0, 3),
                ],
            )
        html = self.client.get("/platform/1").get_data(as_text=True)
        self.assertEqual(html.count("data-gallery-thumb"), 3)
        self.assertLess(html.index("主图占位演示"), html.index("第二张占位演示图"))
        self.assertIn('aria-pressed="true"', html)

    def test_compare_two_three_four_and_validation(self):
        for ids in ("1,5", "1,5,7", "1,5,7,10"):
            response = self.client.get(f"/compare?ids={ids}")
            self.assertEqual(response.status_code, 200)
            self.assertIn(b'data-testid="comparison-table"', response.data)
            positions = [response.data.index(f'data-platform-id="{item}"'.encode()) for item in ids.split(",")]
            self.assertEqual(positions, sorted(positions))
        self.assertEqual(self.client.get("/compare?ids=1,2,3,4,5").status_code, 400)
        self.assertEqual(self.client.get("/compare?ids=1,x").status_code, 400)
        self.assertEqual(self.client.get("/compare?ids=" + "9" * 5000).status_code, 400)
        self.assertEqual(self.client.get("/compare?ids=1").status_code, 200)
        html = self.client.get("/compare?ids=3,5").get_data(as_text=True)
        self.assertIn("—", html)
        self.assertIn("50+ kn（冲刺）", html)

    def test_reimport_preserves_images(self):
        with sqlite3.connect(self.db_path) as db:
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (1,'img/placeholder-platform.svg','保留测试','official',1)"
            )
        import_all(self.db_path)
        with sqlite3.connect(self.db_path) as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platform_images").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT id FROM platforms WHERE name='ALPV / Sea Specter'").fetchone()[0], 1)

    def test_china_append_and_reimports_preserve_all_records_and_images(self):
        append_db = Path(self.temp.name) / "append.sqlite3"
        import_workbook(DEFAULT_XLSX, append_db, source_key=DEFAULT_SOURCE_KEY)
        with sqlite3.connect(append_db) as db:
            old_id = db.execute(
                "SELECT id FROM platforms WHERE name='ALPV / Sea Specter'"
            ).fetchone()[0]
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (?,'img/placeholder-platform.svg','原平台图片','official',1)",
                (old_id,),
            )
            old_platforms = db.execute(
                "SELECT * FROM platforms WHERE country <> '中国' ORDER BY id"
            ).fetchall()
            old_activities = db.execute(
                "SELECT platform_id, sequence_no, description FROM activities ORDER BY platform_id, sequence_no"
            ).fetchall()
            old_sources = db.execute(
                "SELECT platform_id, role, title, url FROM source_links ORDER BY platform_id, role, title, url"
            ).fetchall()

        import_workbook(DEFAULT_CHINA_XLSX, append_db, source_key=CHINA_SOURCE_KEY)
        with sqlite3.connect(append_db) as db:
            china_id = db.execute(
                "SELECT id FROM platforms WHERE name='无人破障艇（公开演练型）'"
            ).fetchone()[0]
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (?,'img/placeholder-platform.svg','中国平台图片','exercise',1)",
                (china_id,),
            )
            self.assertEqual(
                db.execute("SELECT * FROM platforms WHERE country <> '中国' ORDER BY id").fetchall(),
                old_platforms,
            )
            self.assertEqual(
                db.execute(
                    "SELECT platform_id, sequence_no, description FROM activities "
                    "WHERE platform_id <= 16 ORDER BY platform_id, sequence_no"
                ).fetchall(),
                old_activities,
            )
            self.assertEqual(
                db.execute(
                    "SELECT platform_id, role, title, url FROM source_links "
                    "WHERE platform_id <= 16 ORDER BY platform_id, role, title, url"
                ).fetchall(),
                old_sources,
            )
            ids_before = dict(db.execute("SELECT name, id FROM platforms"))
            images_before = db.execute(
                "SELECT platform_id, local_path, caption, image_type, is_primary "
                "FROM platform_images ORDER BY platform_id"
            ).fetchall()

        # Reimport either dataset independently, and then both together. Rows
        # belonging to the other workbook must not be treated as removals.
        import_workbook(DEFAULT_XLSX, append_db, source_key=DEFAULT_SOURCE_KEY)
        import_workbook(DEFAULT_CHINA_XLSX, append_db, source_key=CHINA_SOURCE_KEY)
        import_all(append_db)
        with sqlite3.connect(append_db) as db:
            self.assertEqual(dict(db.execute("SELECT name, id FROM platforms")), ids_before)
            self.assertEqual(
                db.execute(
                    "SELECT platform_id, local_path, caption, image_type, is_primary "
                    "FROM platform_images ORDER BY platform_id"
                ).fetchall(),
                images_before,
            )
            self.assertEqual(db.execute("SELECT COUNT(*) FROM platforms").fetchone()[0], 22)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM activities").fetchone()[0], 30)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM source_links").fetchone()[0], 67)
            self.assertFalse(db.execute(
                "SELECT 1 FROM platforms WHERE source_sheet LIKE '__import_staging__:%' LIMIT 1"
            ).fetchone())

    def test_failed_china_import_rolls_back_without_touching_existing_data(self):
        rollback_db = Path(self.temp.name) / "rollback.sqlite3"
        import_workbook(DEFAULT_XLSX, rollback_db, source_key=DEFAULT_SOURCE_KEY)
        with sqlite3.connect(rollback_db) as db:
            platform_id = db.execute(
                "SELECT id FROM platforms WHERE name='ALPV / Sea Specter'"
            ).fetchone()[0]
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (?,'img/placeholder-platform.svg','回滚保留测试','official',1)",
                (platform_id,),
            )
            # Let the first Chinese row begin importing, then abort on a source
            # from the second row to exercise transaction rollback mid-import.
            db.execute(
                "CREATE TRIGGER abort_china_source BEFORE INSERT ON source_links "
                "WHEN NEW.title LIKE '%瞭望者Ⅱ%' "
                "BEGIN SELECT RAISE(ABORT, 'test rollback'); END"
            )
            before = (
                db.execute("SELECT COUNT(*) FROM platforms").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM activities").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM source_links").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM platform_images").fetchone()[0],
            )

        with self.assertRaises(sqlite3.IntegrityError):
            import_workbook(DEFAULT_CHINA_XLSX, rollback_db, source_key=CHINA_SOURCE_KEY)

        with sqlite3.connect(rollback_db) as db:
            after = (
                db.execute("SELECT COUNT(*) FROM platforms").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM activities").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM source_links").fetchone()[0],
                db.execute("SELECT COUNT(*) FROM platform_images").fetchone()[0],
            )
            self.assertEqual(after, before)
            self.assertEqual(db.execute(
                "SELECT caption FROM platform_images WHERE platform_id=?", (platform_id,)
            ).fetchone()[0], "回滚保留测试")
            self.assertEqual(db.execute(
                "SELECT COUNT(*) FROM platforms WHERE country='中国'"
            ).fetchone()[0], 0)
            self.assertFalse(db.execute(
                "SELECT 1 FROM platforms WHERE source_sheet LIKE '__import_staging__:%' LIMIT 1"
            ).fetchone())

    def test_reordered_workbook_rows_keep_platform_ids_and_images(self):
        reordered = Path(self.temp.name) / "reordered.xlsx"
        workbook = load_workbook(DEFAULT_XLSX, data_only=False)
        sheet = workbook.worksheets[0]
        for column in range(1, sheet.max_column + 1):
            first, second = sheet.cell(4, column), sheet.cell(5, column)
            first_value, second_value = first.value, second.value
            first_link, second_link = copy(first.hyperlink), copy(second.hyperlink)
            first.value, second.value = second_value, first_value
            first._hyperlink, second._hyperlink = second_link, first_link
        workbook.save(reordered)
        with sqlite3.connect(self.db_path) as db:
            before = dict(db.execute("SELECT name, id FROM platforms"))
            db.execute(
                "INSERT INTO platform_images(platform_id,local_path,caption,image_type,is_primary) "
                "VALUES (?,'img/placeholder-platform.svg','reorder test','official',1)",
                (before["ALPV / Sea Specter"],),
            )
        import_workbook(reordered, self.db_path, source_key=DEFAULT_SOURCE_KEY)
        with sqlite3.connect(self.db_path) as db:
            after = dict(db.execute("SELECT name, id FROM platforms"))
            self.assertEqual(before, after)
            self.assertEqual(db.execute(
                "SELECT COUNT(*) FROM platform_images WHERE platform_id=?",
                (before["ALPV / Sea Specter"],),
            ).fetchone()[0], 1)

    def test_migration_is_idempotent_and_preserves_existing_rows(self):
        old_db = Path(self.temp.name) / "old.sqlite3"
        shutil.copy2(self.db_path, old_db)
        with sqlite3.connect(old_db) as db:
            db.execute("DROP TABLE platform_images")
            db.execute("DROP TABLE IF EXISTS schema_migrations")
            before = tuple(db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in ("platforms", "activities", "source_links"))
        self.assertEqual(migrate(old_db), ["001_add_platform_images.sql"])
        self.assertEqual(migrate(old_db), [])
        with sqlite3.connect(old_db) as db:
            after = tuple(db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in ("platforms", "activities", "source_links"))
            self.assertEqual(before, after)
            self.assertTrue(db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='platform_images'").fetchone())

    def test_failed_migration_rolls_back_schema_and_history(self):
        broken_db = Path(self.temp.name) / "broken.sqlite3"
        with sqlite3.connect(broken_db) as db:
            db.execute("CREATE TABLE platforms(id INTEGER PRIMARY KEY)")
        migrations = Path(self.temp.name) / "broken-migrations"
        migrations.mkdir()
        (migrations / "999_broken.sql").write_text(
            "CREATE TABLE leaked(id INTEGER);\nTHIS IS NOT SQL;\n", encoding="utf-8"
        )
        original = migrate_module.MIGRATIONS
        migrate_module.MIGRATIONS = migrations
        try:
            with self.assertRaises(sqlite3.OperationalError):
                migrate_module.migrate(broken_db)
        finally:
            migrate_module.MIGRATIONS = original
        with sqlite3.connect(broken_db) as db:
            self.assertIsNone(db.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='leaked'"
            ).fetchone())
            self.assertEqual(db.execute("SELECT COUNT(*) FROM schema_migrations").fetchone()[0], 0)

    def test_dynamic_statistics(self):
        html = self.client.get("/").get_data(as_text=True)
        self.assertIn("收录平台", html)
        self.assertIn("国家 / 地区", html)
        self.assertIn("活动记录", html)
        self.assertIn(">22<", html)
        self.assertIn(">6<", html)
        self.assertIn(">30<", html)


if __name__ == "__main__":
    unittest.main()
