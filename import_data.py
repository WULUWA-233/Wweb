"""Import the bundled two-batch workbooks into a reproducible SQLite database."""

from __future__ import annotations

import argparse
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_XLSX = ROOT / "data" / "source.xlsx"
DEFAULT_CHINA_XLSX = ROOT / "data" / "source_china.xlsx"
DEFAULT_DB = ROOT / "data" / "platforms.sqlite3"
DEFAULT_SOURCE_KEY = "global"
CHINA_SOURCE_KEY = "china"
BUNDLED_DATASETS = (
    (DEFAULT_SOURCE_KEY, DEFAULT_XLSX),
    (CHINA_SOURCE_KEY, DEFAULT_CHINA_XLSX),
)

LINK_RE = re.compile(r'^=HYPERLINK\("([^"]+)","([^"]+)"\)$', re.IGNORECASE)
SOURCE_KEY_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,63}$")
EVIDENCE_RE = re.compile(r"^\s*(E[2-5])(?=$|[\s+＋,/，；;:：（(])")
PERFORMANCE_FIELDS = (
    "length_text", "beam_draft_text", "weight_displacement_text", "payload_text",
    "speed_text", "range_text", "endurance_text", "propulsion_sea_text",
)
SOURCE_ROLES = {17: "performance", 18: "activity", 19: "control"}
# A dated audit supplement; the original workbook is preserved unchanged.
SOURCE_AUDIT = {
    "Splash Typhoon": {
        "note": "2026-09-22复核：Splash官网首页目前写约350 nmi，原表所引厂商LinkedIn简介写800 nmi；可能是配置或口径差异，暂保留原表数值并标记待核。",
        "link": ("performance", "Splash官网｜当前公开航程", "https://splash9.com/"),
    },
}


@dataclass(frozen=True)
class ParsedDataset:
    source_key: str
    xlsx_path: Path
    source_sheets: tuple[str, ...]
    records: tuple[dict, ...]


def cell_text(cell):
    value = cell.value
    return str(value).strip() if value is not None else None


def source_link(cell):
    """Read HYPERLINK formula or a normal Excel hyperlink; never evaluate formulas."""
    value = cell_text(cell)
    match = LINK_RE.match(value or "")
    if match:
        url, title = match.groups()
    elif cell.hyperlink and cell.hyperlink.target:
        url, title = cell.hyperlink.target, value or cell.hyperlink.target
    else:
        return None
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError(f"Invalid source URL at {cell.coordinate}: {url}")
    return title, url


def evidence_code(raw):
    """Accept only an explicit leading E2-E5 code, never phrases such as 未达E2/非E3."""
    match = EVIDENCE_RE.match(raw or "")
    return match.group(1) if match else None


def activity_parts(raw):
    # Preserve source wording. A semicolon is the workbook's event separator.
    return [part.strip() for part in re.split(r"[；;]", raw or "") if part.strip()]


def source_sheet_name(source_key: str, sheet_title: str) -> str:
    """Return a workbook-scoped sheet identity while preserving legacy primary rows."""
    return sheet_title if source_key == DEFAULT_SOURCE_KEY else f"{source_key}::{sheet_title}"


def sheet_batch(sheet_title: str) -> int:
    if sheet_title.startswith("第一批_"):
        return 1
    if sheet_title.startswith("第二批_"):
        return 2
    raise ValueError(f"Unrecognized batch worksheet: {sheet_title}")


def parse_workbook(xlsx_path=DEFAULT_XLSX, source_key=DEFAULT_SOURCE_KEY) -> ParsedDataset:
    """Validate and parse one workbook without changing SQLite."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.is_file():
        raise FileNotFoundError(xlsx_path)
    if not SOURCE_KEY_RE.fullmatch(source_key or ""):
        raise ValueError("source_key must use 1-64 ASCII letters, digits, dots, underscores or hyphens")

    workbook = load_workbook(xlsx_path, data_only=False, read_only=False)
    try:
        if len(workbook.worksheets) != 2:
            raise ValueError("Expected exactly two source worksheets")

        batches = [sheet_batch(sheet.title) for sheet in workbook.worksheets]
        if sorted(batches) != [1, 2]:
            raise ValueError("Expected one first-batch sheet and one second-batch sheet")

        records = []
        source_sheets = []
        for sheet, batch in zip(workbook.worksheets, batches):
            scoped_sheet = source_sheet_name(source_key, sheet.title)
            source_sheets.append(scoped_sheet)
            for row in range(4, sheet.max_row + 1):
                name = cell_text(sheet.cell(row, 2))
                if not name:
                    continue
                label = cell_text(sheet.cell(row, 3))
                if label not in {"L", "L/W", "W"}:
                    raise ValueError(f"Unexpected label at {sheet.title}!C{row}: {label}")
                values = {
                    "name": name,
                    "country": cell_text(sheet.cell(row, 1)) or "未注明",
                    "batch": batch,
                    "label": label,
                    "platform_type": cell_text(sheet.cell(row, 4)),
                    "payload_description": cell_text(sheet.cell(row, 5)),
                    "control_people": cell_text(sheet.cell(row, 15)),
                    "evidence_raw": cell_text(sheet.cell(row, 16)),
                    "verification_note": cell_text(sheet.cell(row, 20)),
                    "source_sheet": scoped_sheet,
                    "source_row": row,
                    "activity_raw": cell_text(sheet.cell(row, 14)),
                    "links": [],
                }
                values["evidence_code"] = evidence_code(values["evidence_raw"])
                for col, field in enumerate(PERFORMANCE_FIELDS, start=6):
                    values[field] = cell_text(sheet.cell(row, col))
                for col, role in SOURCE_ROLES.items():
                    link = source_link(sheet.cell(row, col))
                    if link:
                        values["links"].append((role, *link))
                audit = SOURCE_AUDIT.get(name)
                if audit:
                    values["verification_note"] = "\n".join(
                        part for part in (values["verification_note"], audit["note"]) if part
                    )
                    values["links"].append(audit["link"])
                records.append(values)
    finally:
        workbook.close()

    if not records or {record["batch"] for record in records} != {1, 2}:
        raise ValueError("Both source worksheets must contain at least one platform")
    incoming_names = [record["name"] for record in records]
    if len(incoming_names) != len(set(incoming_names)):
        raise ValueError("Platform names must be unique across both worksheets")
    return ParsedDataset(source_key, xlsx_path, tuple(source_sheets), tuple(records))


def _import_dataset(db: sqlite3.Connection, dataset: ParsedDataset) -> int:
    """Import one already-parsed dataset inside the caller's transaction."""
    incoming_name_set = {record["name"] for record in dataset.records}
    sheet_placeholders = ",".join("?" for _ in dataset.source_sheets)
    if dataset.source_key == DEFAULT_SOURCE_KEY:
        tracked = db.execute(
            f"SELECT id, name, source_sheet FROM platforms "
            f"WHERE source_sheet IN ({sheet_placeholders})",
            dataset.source_sheets,
        ).fetchall()
    else:
        # SOURCE_KEY_RE excludes GLOB metacharacters. Tracking the whole
        # namespace also catches a removed or renamed worksheet.
        tracked = db.execute(
            "SELECT id, name, source_sheet FROM platforms WHERE source_sheet GLOB ?",
            (f"{dataset.source_key}::*",),
        ).fetchall()

    removed_or_renamed = sorted(row[1] for row in tracked if row[1] not in incoming_name_set)
    if removed_or_renamed:
        raise ValueError(
            f"Dataset {dataset.source_key!r} removed or renamed existing platforms; "
            "resolve explicitly before import: " + ", ".join(removed_or_renamed)
        )

    tracked_ids = {row[0] for row in tracked}
    name_placeholders = ",".join("?" for _ in incoming_name_set)
    existing_names = db.execute(
        f"SELECT id, name, source_sheet FROM platforms WHERE name IN ({name_placeholders})",
        tuple(sorted(incoming_name_set)),
    ).fetchall()
    foreign_conflicts = sorted(row[1] for row in existing_names if row[0] not in tracked_ids)
    if foreign_conflicts:
        raise ValueError(
            f"Dataset {dataset.source_key!r} conflicts with platform names owned by another "
            "source: " + ", ".join(foreign_conflicts)
        )

    # Temporarily free the unique (source_sheet, source_row) coordinates. This
    # permits row insertion/reordering while stable names retain ids and images.
    for platform_id, _name, _sheet in tracked:
        db.execute(
            "UPDATE platforms SET source_sheet = ?, source_row = ? WHERE id = ?",
            (f"__import_staging__:{dataset.source_key}:{platform_id}", -platform_id, platform_id),
        )

    columns = (
        "name", "country", "batch", "label", "platform_type", "payload_description",
        *PERFORMANCE_FIELDS, "control_people", "evidence_code", "evidence_raw",
        "verification_note", "source_sheet", "source_row",
    )
    placeholders = ", ".join("?" for _ in columns)
    updates = ", ".join(f"{column}=excluded.{column}" for column in columns[1:])
    insert = (
        f"INSERT INTO platforms ({', '.join(columns)}) VALUES ({placeholders}) "
        f"ON CONFLICT(name) DO UPDATE SET {updates}"
    )
    for record in dataset.records:
        db.execute(insert, [record.get(key) for key in columns])
        platform_id = db.execute(
            "SELECT id FROM platforms WHERE name = ?", (record["name"],)
        ).fetchone()[0]
        # Refresh workbook-owned child rows only. User-maintained images survive.
        db.execute("DELETE FROM activities WHERE platform_id = ?", (platform_id,))
        db.execute("DELETE FROM source_links WHERE platform_id = ?", (platform_id,))
        for sequence_no, description in enumerate(activity_parts(record["activity_raw"]), start=1):
            db.execute(
                "INSERT INTO activities (platform_id, sequence_no, description) VALUES (?, ?, ?)",
                (platform_id, sequence_no, description),
            )
        for role, title, url in record["links"]:
            db.execute(
                "INSERT OR IGNORE INTO source_links (platform_id, role, title, url) VALUES (?, ?, ?, ?)",
                (platform_id, role, title, url),
            )

    if tracked_ids:
        id_placeholders = ",".join("?" for _ in tracked_ids)
        if db.execute(
            f"SELECT 1 FROM platforms WHERE id IN ({id_placeholders}) "
            "AND source_sheet GLOB '__import_staging__:*' LIMIT 1",
            tuple(sorted(tracked_ids)),
        ).fetchone():
            raise RuntimeError("Import staging rows remain; transaction has been rolled back")
    return len(dataset.records)


def import_datasets(datasets, db_path=DEFAULT_DB) -> int:
    """Parse datasets first, then import all of them in one SQLite transaction."""
    parsed = [
        item if isinstance(item, ParsedDataset) else parse_workbook(item[1], item[0])
        for item in datasets
    ]
    source_keys = [item.source_key for item in parsed]
    if len(source_keys) != len(set(source_keys)):
        raise ValueError("Each dataset source_key must be unique")
    all_names = [record["name"] for item in parsed for record in item.records]
    if len(all_names) != len(set(all_names)):
        raise ValueError("Platform names must be unique across all imported datasets")

    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(db_path)
    try:
        db.execute("PRAGMA foreign_keys = ON")
        db.executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
        with db:
            count = sum(_import_dataset(db, item) for item in parsed)
    finally:
        db.close()
    return count


def import_workbook(xlsx_path=DEFAULT_XLSX, db_path=DEFAULT_DB, source_key=DEFAULT_SOURCE_KEY):
    """Import one workbook. Use a unique source_key for every additional workbook."""
    return import_datasets(((source_key, Path(xlsx_path)),), db_path)


def import_all(db_path=DEFAULT_DB):
    """Import every workbook bundled with this project atomically."""
    return import_datasets(BUNDLED_DATASETS, db_path)


def infer_bundled_source_key(xlsx_path: Path):
    resolved = xlsx_path.resolve()
    for source_key, path in BUNDLED_DATASETS:
        if resolved == path.resolve():
            return source_key
    return None


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="导入海上无人平台数据")
    parser.add_argument(
        "--xlsx", type=Path,
        help="只导入一个 Excel 文件；省略时原子化导入 data 下的全部内置工作簿",
    )
    parser.add_argument(
        "--source-key",
        help="自定义工作簿的稳定来源键（多个工作簿之间必须唯一）",
    )
    parser.add_argument("--db", type=Path, default=DEFAULT_DB, help="SQLite 输出路径")
    args = parser.parse_args()
    if args.xlsx:
        source_key = args.source_key or infer_bundled_source_key(args.xlsx)
        if not source_key:
            parser.error("自定义 --xlsx 必须同时提供唯一的 --source-key")
        count = import_workbook(args.xlsx, args.db, source_key)
        scope = f"来源 {source_key}"
    else:
        if args.source_key:
            parser.error("--source-key 只能与 --xlsx 同时使用")
        count = import_all(args.db)
        scope = "全部内置工作簿"
    print(f"已从{scope}导入 {count} 型平台：{args.db}")
